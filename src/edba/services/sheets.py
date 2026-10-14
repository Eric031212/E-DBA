"""Read small spreadsheets (.xlsx / .csv) into header + rows of strings."""

from __future__ import annotations

import csv
import io
from pathlib import Path

from edba.errors import DomainError

MAX_ROWS = 5000


def _cell(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_table(filename: str, data: bytes) -> tuple[list[str], list[list[str]]]:
    ext = Path(filename).suffix.lower()
    if ext == ".csv":
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as e:
            raise DomainError("CSV files must be UTF-8 encoded.") from e
        raw = [list(r) for r in csv.reader(io.StringIO(text))]
    elif ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:
            raise DomainError("Could not read the spreadsheet; is it a valid .xlsx file?") from e
        ws = wb.worksheets[0]
        raw = [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]
    else:
        raise DomainError("Upload an Excel (.xlsx) or CSV file.")
    rows = [[_cell(c) for c in r] for r in raw]
    rows = [r for r in rows if any(r)]
    if not rows:
        raise DomainError("The file is empty.")
    header = [h.strip().lower().replace(" ", "_") for h in rows[0]]
    while header and not header[-1]:
        header.pop()
    body = [(r + [""] * len(header))[: len(header)] for r in rows[1:]]
    if len(body) > MAX_ROWS:
        raise DomainError(f"At most {MAX_ROWS} rows per upload.")
    return header, body
